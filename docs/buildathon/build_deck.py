"""Build the PlutusShield Midnight Buildathon (Wave 2) deck.

    python3 -m venv .venv && .venv/bin/pip install python-pptx
    .venv/bin/python docs/buildathon/build_deck.py
    soffice --headless --convert-to pdf docs/buildathon/PlutusShield-Wave2-deck.pptx --outdir docs/buildathon

Every number and hash here comes from the repo's deployment records
(contracts/*/deployments) and the test run; testnet only, not audited.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

HERE = Path(__file__).parent
BG = RGBColor(0x0B, 0x0D, 0x1A)
CARD = RGBColor(0x15, 0x18, 0x2B)
LINE = RGBColor(0x2A, 0x2F, 0x4A)
INK = RGBColor(0xE9, 0xEB, 0xF5)
MUTED = RGBColor(0x9A, 0xA0, 0xBD)
BLUE = RGBColor(0x7C, 0x9C, 0xFF)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
GOLD = RGBColor(0xF2, 0xC9, 0x4C)
SERIF = "Georgia"
SANS = "Helvetica"
MONO = "Courier New"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(s, x, y, w, h, runs, size=18, color=INK, font=SANS, align=PP_ALIGN.LEFT, spacing=1.1, after=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, color, bold, font)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(size * 0.45 if after is None else after)
        parts = [(para, color, False, font)] if isinstance(para, str) else para
        for t, c, b, f in parts:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(size)
            r.font.color.rgb = c or color
            r.font.bold = b
            r.font.name = f or font
    return tb


def kicker(s, label, y=0.55):
    text(s, 0.8, y, 10, 0.4, label.upper(), size=12, color=GOLD, font=MONO)


def title(s, t, y=0.95, size=40):
    text(s, 0.8, y, 11.7, 1.2, t, size=size, font=SERIF)


def card(s, x, y, w, h, head, body, head_color=BLUE, size=15):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = CARD
    sh.line.color.rgb = LINE; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    text(s, x + 0.3, y + 0.25, w - 0.6, 0.5, head, size=size + 3, color=head_color, font=SANS)
    text(s, x + 0.3, y + 0.8, w - 0.6, h - 1.0, body, size=size, color=MUTED)


def footer(s, n):
    text(s, 0.8, 6.95, 8, 0.3, "PlutusShield · Midnight Buildathon Wave 2 · testnet only, not audited",
         size=10, color=MUTED, font=MONO)
    text(s, 11.5, 6.95, 1.0, 0.3, f"{n:02d}", size=10, color=MUTED, font=MONO, align=PP_ALIGN.RIGHT)


def shot(s, name, x, y, w):
    s.shapes.add_picture(str(HERE / "shots" / f"{name}.png"), Inches(x), Inches(y), width=Inches(w))


n = 0
# 1 Title
s = slide(); n += 1
kicker(s, "Midnight Buildathon · Wave 2", y=1.2)
text(s, 0.8, 1.6, 7.0, 3.2, [[("Protection you", INK, False, SERIF)], [("can ", INK, False, SERIF), ("prove.", BLUE, False, SERIF)],
                            [("Positions you", INK, False, SERIF)], [("never ", INK, False, SERIF), ("expose.", VIOLET, False, SERIF)]],
     size=44, spacing=0.95, after=0)
text(s, 0.8, 4.75, 6.6, 1.4, "PlutusShield is DeFi and smart-contract insurance. Cardano holds the money and pays "
     "claims. Midnight keeps who is covered, for how much, and the claim evidence private.", size=17, color=MUTED)
text(s, 0.8, 6.2, 7, 0.4, "kshot3000.github.io/PlutusShield  ·  github.com/Kshot3000/PlutusShield",
     size=13, color=BLUE, font=MONO)
shot(s, "home", 7.4, 1.9, 5.4)
footer(s, n)

# 2 Problem
s = slide(); n += 1
kicker(s, "The problem"); title(s, "DeFi cover today makes you choose: get paid, or stay private.")
card(s, 0.8, 2.6, 3.8, 3.6, "Cover leaks your book",
     "On a public chain a policy says who you are, what you hold and how much you insured. "
     "That is a map for anyone hunting large positions.")
card(s, 4.77, 2.6, 3.8, 3.6, "Claims leak your evidence",
     "Exploit claims need proof: tx traces, screenshots, wallet history. Posting it publicly "
     "exposes the claimant and sometimes the bug.")
card(s, 8.74, 2.6, 3.8, 3.6, "Payouts need trust",
     "Most cover relies on one assessor or a token vote. Users need to see that no single key "
     "can approve or block a payout.")
footer(s, n)

# 3 How it works
s = slide(); n += 1
kicker(s, "How it works"); title(s, "Cardano settles. Midnight keeps the secrets.")
card(s, 0.8, 2.4, 5.7, 3.9, "Cardano Preview · Aiken, Plutus V3", [
    "Premiums paid in ADA or test USDC (tUSDCx) into on-chain pools.",
    "Depeg cover pays from oracle feeds; sales close by themselves if the feeds go stale or the peg breaks.",
    "Exploit cover pays only with 2-of-3 committee signatures, enforced by the validator.",
    "The policy datum holds a hash commitment, not the terms."], head_color=BLUE)
card(s, 6.83, 2.4, 5.7, 3.9, "Midnight Preprod · Compact, ZK", [
    "registerPolicy stores a commitment to the holder and terms.",
    "proveCover(policyId, minCoverage) proves cover is at least X without revealing the real limit.",
    "fileClaim puts only a commitment to the encrypted evidence on chain.",
    "voteClaim tallies three committee seats; 2 approvals mark the claim PAID."], head_color=VIOLET)
footer(s, n)

# 4 Architecture
s = slide(); n += 1
kicker(s, "Architecture"); title(s, "One policy, two chains, a relay in between.")
arch = """ Cardano Preview (Aiken, Plutus V3)            Midnight Preprod (Compact, ZK)
 ─────────────────────────────────             ───────────────────────────────
 Buy tx: premium + policy datum   ──relay──►   registerPolicy(commitment)
   (holds only a hash commitment)              proveCover(id, ≥min)  [ZK, holder]
 Depeg pool: oracle feeds (2h)                 fileClaim(id, evidenceCommitment)
 Exploit pool: Settle needs        ◄─decision─ voteClaim × 3 seats → PAID at 2-of-3
   2-of-3 committee signatures                 (claim round resets on reject)
 Payout in ADA / tUSDCx                        Evidence: encrypted off chain, hash on chain"""
sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.3), Inches(11.73), Inches(3.4))
sh.adjustments[0] = 0.04; sh.fill.solid(); sh.fill.fore_color.rgb = CARD; sh.line.color.rgb = LINE
text(s, 1.05, 2.5, 11.3, 3.1, arch.split("\n"), size=14, color=INK, font=MONO, spacing=1.0)
text(s, 0.8, 5.9, 11.7, 0.9, "No wallet address, position size or evidence lands on Midnight in the clear. "
     "The relay only moves commitments one way and the committee decision the other.", size=16, color=MUTED)
footer(s, n)

# 5 Midnight contracts
s = slide(); n += 1
kicker(s, "Built on Midnight"); title(s, "Two Compact contracts, deployed and exercised on Preprod.")
card(s, 0.8, 2.4, 5.7, 3.9, "policy-cover v1", [
    [("Preprod contract  ", MUTED, False, SANS), ("d84c6187…925e8787", INK, False, MONO)],
    "Circuits: registerPolicy, proveCover, rotateHolder, fileClaim, resolveClaim, expirePolicy.",
    "5 policies registered and proven, 4 claims filed and resolved.",
    "Single assessor: the first version, kept for the record."], head_color=VIOLET)
card(s, 6.83, 2.4, 5.7, 3.9, "policy-cover v2", [
    [("Preprod contract  ", MUTED, False, SANS), ("5430b803…4feb8110", INK, False, MONO)],
    "Adds voteClaim: a 2-of-3 committee with per-round tallies; a reject resets the round.",
    "First drill: approve, reject, approve, then PAID at 2 of 3.",
    "Mirrors all 3 mirrorable Cardano Preview policies."], head_color=VIOLET)
text(s, 0.8, 6.4, 11.7, 0.4, "compactc 0.31.1 · both contracts full-compile with prover and verifier keys",
     size=13, color=MUTED, font=MONO)
footer(s, n)

# 6 Proof points
s = slide(); n += 1
kicker(s, "Live on testnet"); title(s, "Every line here is backed by a transaction.")
rows = [
    ("Depeg pool (ADA + tUSDCx)", "11 live Preview policies, buy / claim / release from a CIP-30 wallet"),
    ("First exploit claim paid", "50 tADA: Midnight claim PAID, then Cardano Settle 5509c2d9…"),
    ("First committee payout", "50 tADA co-signed by 2 of 3 assessors, Settle c35535dd…"),
    ("1-of-3 payout attempt", "Rejected by the validator in both local and Blockfrost evaluation"),
    ("Depeg oracle", "45 recorded live publishes to 3 Preview feeds, latest 99.84% (Oct 12)"),
    ("Midnight relay", "registerPolicy + proveCover on Preprod for 5 policies, readable via Blockfrost"),
]
for i, (a, b) in enumerate(rows):
    y = 2.35 + i * 0.72
    text(s, 0.8, y, 4.6, 0.6, a, size=17, color=INK)
    text(s, 5.5, y, 7.1, 0.6, b, size=15, color=MUTED)
footer(s, n)

# 7 Product
s = slide(); n += 1
kicker(s, "The product"); title(s, "A real app: quote, buy, prove, claim.")
shot(s, "cover", 0.8, 2.3, 5.75)
shot(s, "evidence", 6.78, 2.3, 5.75)
text(s, 0.8, 5.65, 5.75, 0.8, "/cover: connect a Preview wallet, price depeg or exploit cover, pay in ADA or tUSDCx.",
     size=14, color=MUTED)
text(s, 6.78, 5.65, 5.75, 0.8, "/claim/evidence: evidence is encrypted in the browser (AES-256-GCM); only its "
     "SHA-256 commitment goes to fileClaim.", size=14, color=MUTED)
footer(s, n)

# 8 Engineering
s = slide(); n += 1
kicker(s, "Engineering"); title(s, "305 unit tests plus an end-to-end emulator run, all green.")
stats = [("131", "Aiken validator tests"), ("97", "SDK: quote engine, codecs, evidence vault, claim kit, relay"),
         ("60", "Compact circuit simulation, claims relay, relay plan, state decode"), ("14 + 3", "oracle relay · committee kit")]
for i, (big, small) in enumerate(stats):
    x = 0.8 + i * 3.0
    text(s, x, 2.6, 2.8, 1.0, big, size=48, color=BLUE, font=SERIF)
    text(s, x, 3.7, 2.7, 1.2, small, size=15, color=MUTED)
text(s, 0.8, 5.1, 11.7, 1.4, [
    "Lucid emulator e2e: pool, buy, settle, expire, exploit 2-of-3, and every 1-of-3 combination rejected.",
    "Next.js static site, 17 routes, redeployed to GitHub Pages on every push; checked on iPhone (WebKit) and Pixel."],
    size=16, color=INK)
footer(s, n)

# 9 Honest + next
s = slide(); n += 1
kicker(s, "Where it stands"); title(s, "Honest about today. Clear about next.")
card(s, 0.8, 2.4, 5.7, 3.9, "Today", [
    "Testnet only (Cardano Preview, Midnight Preprod). Not audited.",
    "The three committee seats are run by the team.",
    "Exploit evidence in the drills is demo evidence.",
    "Midnight retired its hosted Preprod indexer; reads now go through Blockfrost."], head_color=GOLD)
card(s, 6.83, 2.4, 5.7, 3.9, "Next", [
    "Independent assessor operators (offline keygen and signed enrollment kit is in the repo).",
    "proveCover and fileClaim straight from the browser wallet (claims already file from the holder's own claim kit).",
    "Security audit, then a capped mainnet launch.",
    "Catalyst funding for claim capital."], head_color=BLUE)
footer(s, n)

# 10 Close
s = slide(); n += 1
kicker(s, "Try it", y=1.6)
text(s, 0.8, 2.1, 11.7, 1.4, [[("Insurance that pays on Cardano and ", INK, False, SERIF), ("stays quiet", VIOLET, False, SERIF),
                              (" on Midnight.", INK, False, SERIF)]], size=40)
text(s, 0.8, 3.9, 11.7, 1.8, [
    [("Live app   ", MUTED, False, SANS), ("https://kshot3000.github.io/PlutusShield/", BLUE, False, MONO)],
    [("Code       ", MUTED, False, SANS), ("https://github.com/Kshot3000/PlutusShield", BLUE, False, MONO)],
    [("Built by   ", MUTED, False, SANS), ("@kshot9000", BLUE, False, MONO)]], size=18)
footer(s, n)

out = HERE / "PlutusShield-Wave2-deck.pptx"
prs.save(out)
print("wrote", out)
